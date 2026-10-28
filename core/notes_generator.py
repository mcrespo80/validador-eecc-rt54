import os
import re
from typing import Dict, Any, List, Tuple, Optional
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

from core.extractor import extract_numbers_from_line
from core.report_pdf import FONT_REG, FONT_BOLD, FONT_ITA, NAVY_PRIMARY, TEXT_DARK, TEXT_MUTED, LINE_COLOR, BG_LIGHT_ROW, WHITE
from core.remediator import ConsejoSequentialCanvas
from core.validator import fmt_ars


# Mapping of rubric patterns (both RT 54 NUA and legacy names) to official CPCECABA RT 54 policy templates
RUBRO_POLICY_CATALOG: List[Dict[str, Any]] = [
    {
        "id": "caja_bancos",
        "order": 1,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 207-212",
        "patterns": [r"^Caja\s+y\s+bancos\b", r"^Disponibilidades\b"],
        "default_name": "Caja y bancos",
        "policy_ep": "El efectivo disponible y los saldos en cuentas bancarias en pesos se encuentran medidos a su valor nominal. Los saldos en moneda extranjera, de existir, han sido convertidos al tipo de cambio vigente al cierre del ejercicio.",
        "policy_re": "El efectivo disponible y los saldos en cuentas bancarias en pesos se encuentran medidos a su valor nominal. Los saldos en moneda extranjera, de existir, han sido convertidos al tipo de cambio vigente al cierre del ejercicio."
    },
    {
        "id": "inversiones_financieras",
        "order": 2,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 213-233",
        "patterns": [r"^Inversiones\s+financieras\b", r"^Inversiones\s+temporarias\b", r"^Inversiones\s+corrientes\b", r"^Inversiones$"],
        "default_name": "Inversiones financieras",
        "policy_ep": "Las colocaciones en depósitos a plazo fijo y similares han sido medidas a su costo amortizado al cierre del ejercicio. Las tenencias en fondos comunes de inversión y títulos con cotización han sido medidas a su valor razonable sobre la base del valor de cotización o cuotaparte vigente a la fecha de cierre. El valor registrado de las inversiones financieras medidas a costo amortizado no supera su valor recuperable.",
        "policy_re": "Las colocaciones en depósitos a plazo fijo y títulos mantenidos hasta su vencimiento han sido medidas a su costo amortizado utilizando el método de la tasa de interés efectiva. Los fondos comunes de inversión y títulos públicos con intención de negociación han sido medidos a su valor razonable en base a su cotización vigente al cierre de cada ejercicio. El valor registrado no supera su valor recuperable."
    },
    {
        "id": "cuentas_por_cobrar",
        "order": 3,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 234-265",
        "patterns": [
            r"^Cuentas\s+por\s+cobrar\s+a\s+clientes(?:\s+en\s+moneda)?\b",
            r"^Cuentas\s+por\s+cobrar\s+a\s+asociados\b",
            r"^Cuentas\s+por\s+cobrar\s+a\s+terceros\b",
            r"^Cuentas\s+por\s+cobrar\s+por\s+prestaciones\b",
            r"^Cr[eé]ditos\s+por\s+ventas\b",
            r"^Deudores\s+por\s+ventas\b",
            r"^Cr[eé]ditos\s+por\s+servicios\b",
            r"^Deudores\s+por\s+aportes\s+y\s+contribuciones\b"
        ],
        "default_name": "Cuentas por cobrar a clientes en moneda / Créditos por ventas",
        "policy_ep": "Se miden a su valor nominal (haciendo uso de la simplificación para Entidades Pequeñas del párrafo 247 de la RT 54 cuando no existen componentes financieros explícitos), netos de la previsión por desvalorización e incobrabilidad, de corresponder. A cada fecha de cierre se evalúa la existencia de indicios de desvalorización en función de la antigüedad de la cartera y la evolución de las cobranzas.",
        "policy_re": "Las partidas sobre las que se segregaron componentes financieros implícitos o que devengan intereses explícitos se miden a su costo amortizado utilizando el método de la tasa de interés efectiva; en los restantes casos de corto plazo se miden a su valor nominal, netos de la previsión por desvalorización e incobrabilidad calculada en función del riesgo de cartera."
    },
    {
        "id": "creditos_impositivos",
        "order": 4,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 266-280",
        "patterns": [r"^Cr[eé]ditos\s+impositivos\b", r"^Cr[eé]ditos\s+fiscales\b"],
        "default_name": "Créditos impositivos",
        "policy_ep": "Los saldos a favor, anticipos, retenciones y percepciones impositivas se miden a su valor nominal, netos de previsión por desvalorización cuando existan dudas sobre su recuperabilidad o utilización futura contra obligaciones fiscales proyectadas.",
        "policy_re": "Los saldos a favor, anticipos, retenciones y percepciones impositivas se miden a su valor nominal (o costo amortizado cuando su realización supere los doce meses), evaluando su recuperabilidad en función de las proyecciones fiscales futuras."
    },
    {
        "id": "otras_cuentas_por_cobrar",
        "order": 5,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 266-289",
        "patterns": [
            r"^Otras\s+cuentas\s+por\s+cobrar(?:\s+en\s+moneda)?\b",
            r"^Otros\s+cr[eé]ditos(?:\s+en\s+moneda)?\b",
            r"^Cr[eé]ditos\s+en\s+moneda\s+con\s+partes\s+relacionadas\b"
        ],
        "default_name": "Otras cuentas por cobrar en moneda / Otros créditos",
        "policy_ep": "Se miden a su valor nominal reexpresado (para partidas no monetarias o anticipos) o al importe nominal a cobrar en moneda de cierre, netos de previsión por desvalorización de corresponder.",
        "policy_re": "Se miden a su costo amortizado cuando existen plazos de cobro diferidos superiores a doce meses o intereses explícitos, y a su valor nominal en los restantes casos, netos de previsión por desvalorización."
    },
    {
        "id": "bienes_cambio",
        "order": 6,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 290-322",
        "patterns": [
            r"^Bienes\s+de\s+cambio\b",
            r"^Bienes\s+para\s+consumo\s+o\s+comercializaci[oó]n\b",
            r"^Inventarios\b"
        ],
        "default_name": "Bienes de cambio / Bienes para consumo o comercialización",
        "policy_ep": "Se miden a su costo de última compra o costo de reposición/adquisición reexpresado en moneda homogénea de cierre de acuerdo con lo indicado en la Nota 1.3 (conforme a las opciones de los párrafos 291 a 295 de la RT 54). El valor registrado no supera su valor recuperable estimado al cierre del ejercicio.",
        "policy_re": "Los bienes adquiridos se miden a su costo de reposición o costo de la última compra reexpresado a moneda de cierre, y los productos en proceso o terminados a su costo de producción reexpresado. El valor contable no supera su valor neto de realización ni su valor recuperable."
    },
    {
        "id": "bienes_uso",
        "order": 7,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 323-358",
        "patterns": [r"^Bienes\s+de\s+uso\b", r"^Propiedad,\s+planta\s+y\s+equipo\b"],
        "default_name": "Bienes de uso",
        "policy_ep": "Están medidos a su costo de adquisición reexpresado en moneda homogénea de cierre de acuerdo con lo indicado en la Nota 1.3, neto de sus correspondientes depreciaciones acumuladas. Las depreciaciones han sido calculadas por el método de la línea recta sobre los valores reexpresados, aplicando tasas anuales suficientes para extinguir sus valores al final de la vida útil estimada de los bienes. El valor contable de los bienes de uso no supera su valor recuperable.",
        "policy_re": "Están medidos a su costo de adquisición reexpresado en moneda homogénea de cierre de acuerdo con lo indicado en la Nota 1.3, neto de sus depreciaciones acumuladas calculadas por el método de la línea recta en función de la vida útil económica estimada para cada grupo de bienes. El valor registrado no supera su valor recuperable."
    },
    {
        "id": "propiedades_inversion",
        "order": 8,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 359-374",
        "patterns": [r"^Propiedades\s+de\s+inversi[oó]n\b"],
        "default_name": "Propiedades de inversión",
        "policy_ep": "Se encuentran medidas bajo el modelo de costo de adquisición reexpresado de acuerdo con la Nota 1.3, neto de sus depreciaciones acumuladas calculadas por el método de la línea recta sobre la vida útil estimada de los inmuebles (RT 54 párr. 362).",
        "policy_re": "Se encuentran medidas a su costo de adquisición reexpresado de acuerdo con la Nota 1.3, neto de sus depreciaciones acumuladas calculadas por el método de la línea recta. Su valor razonable no presenta desvalorizaciones respecto del valor contable."
    },
    {
        "id": "activos_intangibles",
        "order": 9,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 375-410",
        "patterns": [r"^Activos\s+intangibles\b", r"^Bienes\s+intangibles\b"],
        "default_name": "Activos intangibles",
        "policy_ep": "Están medidos a su costo de adquisición o desarrollo reexpresado en moneda de cierre según la Nota 1.3, neto de sus amortizaciones acumuladas calculadas por el método de la línea recta en función de su vida útil estimada.",
        "policy_re": "Están medidos a su costo de adquisición o desarrollo reexpresado en moneda de cierre según la Nota 1.3, neto de sus amortizaciones acumuladas calculadas por el método de la línea recta."
    },
    {
        "id": "otras_inversiones_vpp",
        "order": 10,
        "section": "Activo",
        "norm_ref": "RT 54 párr. 411-445",
        "patterns": [r"^Participaciones\s+en\s+subsidiarias", r"^Otras\s+inversiones\b", r"^Inversiones\s+permanentes\b"],
        "default_name": "Otras inversiones / Participaciones permanentes",
        "policy_ep": "Las otras inversiones cuyos activos cuentan con precio directamente observable se miden a su valor razonable; en los restantes casos se miden al costo de adquisición reexpresado en moneda homogénea de cierre según la Nota 1.3 (o por el método del Valor Patrimonial Proporcional cuando existe control o influencia significativa).",
        "policy_re": "Las participaciones en subsidiarias, asociadas y negocios conjuntos societarios se miden aplicando el método del Valor Patrimonial Proporcional (VPP) sobre los últimos estados contables disponibles al cierre, ajustados en moneda homogénea."
    },
    {
        "id": "proveedores_deudas_comerciales",
        "order": 11,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 458-472",
        "patterns": [
            r"^Proveedores\s+de\s+bienes\s+y\s+servicios\b",
            r"^Deudas\s+con\s+proveedores\s+de\s+bienes\s+y\s+servicios\b",
            r"^Deudas\s+comerciales\b",
            r"^Cuentas\s+por\s+pagar\b",
            r"^Prestadores\s+y\s+proveedores\b"
        ],
        "default_name": "Proveedores de bienes y servicios / Deudas comerciales",
        "policy_ep": "Se miden a su valor nominal al cierre del ejercicio (aplicando la simplificación para Entidades Pequeñas prevista en el párrafo 464 de la RT 54).",
        "policy_re": "Se miden a su costo amortizado cuando se han segregado componentes financieros implícitos en operaciones a plazo superior a doce meses o con intereses explícitos, y a su valor nominal en los restantes casos de corto plazo."
    },
    {
        "id": "prestamos_pasivos_financieros",
        "order": 12,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 458-472",
        "patterns": [
            r"^Pr[eé]stamos\s+y\s+otros\s+pasivos\s+financieros\b",
            r"^Deudas\s+bancarias(?:\s+y\s+financieras)?\b",
            r"^Deudas\s+financieras\b",
            r"^Pr[eé]stamos\b"
        ],
        "default_name": "Préstamos y otros pasivos financieros",
        "policy_ep": "Los préstamos bancarios y otros pasivos financieros que devengan intereses explícitos se miden a su costo amortizado utilizando el método de la tasa de interés efectiva, incluyendo los intereses devengados hasta la fecha de cierre del ejercicio (RT 54 párr. 464).",
        "policy_re": "Los préstamos y pasivos financieros se miden a su costo amortizado utilizando el método de la tasa de interés efectiva, reconociendo los costos financieros devengados hasta el cierre."
    },
    {
        "id": "deudas_laborales",
        "order": 13,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 473-480",
        "patterns": [
            r"^Deudas\s+laborales\s+y\s+previsionales\b",
            r"^Remuneraciones\s+y\s+cargas\s+sociales\b",
            r"^Deudas\s+sociales\b",
            r"^Deudas\s+previsionales\b"
        ],
        "default_name": "Deudas laborales y previsionales",
        "policy_ep": "Comprenden las remuneraciones, sueldo anual complementario devengado, provisiones para vacaciones y las contribuciones y aportes al régimen de seguridad social y sindicales pendientes de ingreso al cierre, medidos a su valor nominal.",
        "policy_re": "Comprenden las obligaciones laborales y de la seguridad social devengadas al cierre (incluyendo provisiones de SAC y vacaciones), medidas a su valor nominal (o costo amortizado en caso de planes de pago a largo plazo)."
    },
    {
        "id": "deudas_fiscales",
        "order": 14,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 481-485",
        "patterns": [
            r"^Deudas\s+fiscales\b",
            r"^Cargas\s+fiscales\b",
            r"^Deudas\s+impositivas\b"
        ],
        "default_name": "Deudas fiscales",
        "policy_ep": "Comprenden las obligaciones tributarias nacionales, provinciales y municipales devengadas al cierre del ejercicio, netas de pagos a cuenta, medidas a su valor nominal (o costo amortizado para planes de facilidades de pago con intereses explícitos).",
        "policy_re": "Comprenden las obligaciones tributarias devengadas al cierre del ejercicio medidas a su valor nominal, y los planes de facilidades de pago medidos a su costo amortizado."
    },
    {
        "id": "otras_deudas_anticipos_fondos",
        "order": 15,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 486-492 / 731",
        "patterns": [
            r"^Otras\s+deudas(?:\s+en\s+moneda)?\b",
            r"^Otros\s+pasivos\b",
            r"^Deudas\s+en\s+especie\b",
            r"^Anticipos\s+de\s+clientes\b",
            r"^Subsidios\s+y\s+otras\s+ayudas\s+gubernamentales\b",
            r"^Fondos\s+con\s+destino\s+espec[ií]fico\b",
            r"^Deudas\s+en\s+moneda\s+con\s+partes\s+relacionadas\b"
        ],
        "default_name": "Otras deudas / Otros pasivos",
        "policy_ep": "Las otras deudas en moneda y los fondos con destino específico (RT 54 párr. 731) se miden a su valor nominal. Los anticipos o deudas en especie, de existir, se miden conforme al importe reexpresado recibido o al costo de cancelación.",
        "policy_re": "Se miden a su valor nominal en el corto plazo o costo amortizado cuando superan los doce meses. Las obligaciones en especie se miden por el mayor valor entre la medición inicial reexpresada del anticipo recibido y el costo inevitable de cumplimiento."
    },
    {
        "id": "previsiones_pasivo",
        "order": 16,
        "section": "Pasivo",
        "norm_ref": "RT 54 párr. 493-507",
        "patterns": [r"^Previsiones$|^Previsiones\s+para\b"],
        "default_name": "Previsiones",
        "policy_ep": "Las previsiones incluidas en el pasivo corresponden a contingencias probables al cierre susceptibles de cuantificación objetiva y se miden al valor nominal de la mejor estimación de los montos a desembolsar a la fecha de cierre (Nota 2.13 Modelo CPCECABA), considerando la opinión de los asesores legales.",
        "policy_re": "Las previsiones incluidas en el pasivo se miden al valor descontado de los flujos de efectivo esperados de la mejor estimación de los montos a cada fecha de cierre (Nota 2.15 Modelo CPCECABA), tomando en cuenta las expectativas de la Dirección y el informe de los asesores legales."
    }
]


LEGACY_RUBRO_KEYWORDS = [
    "créditos por ventas",
    "creditos por ventas",
    "deudores por ventas",
    "deudas comerciales",
    "remuneraciones y cargas",
    "deudas sociales",
    "cargas fiscales",
    "otros créditos",
    "otros creditos",
    "disponibilidades"
]


def detect_active_rubros_from_balance(lines: List[Dict[str, Any]], full_text: str) -> List[Dict[str, Any]]:
    """
    Scans the uploaded balance (ESP, ER/ERG, and Cuadros de Notas Activo/Pasivo) and extracts
    the exact rubric names used by the entity that have a non-zero balance in Actual or Comparativo
    (or all rubrics present if the user uploaded a structure without zero rows).
    Preserves the exact rubric denomination written in the user's balance!
    """
    found_rubros: List[Dict[str, Any]] = []
    seen_ids = set()

    for item in lines:
        txt = item["text"].strip()
        if not txt:
            continue

        # Split side-by-side table rows (Activo | Pasivo)
        segments = [s.strip() for s in txt.split("|")]
        nums = extract_numbers_from_line(txt)

        # Determine if this row is explicitly zero in both columns
        is_explicit_zero = False
        if len(nums) >= 2 and all(abs(x) < 0.005 for x in nums[:2]):
            is_explicit_zero = True
        elif re.search(r"(?:\|\s*-\s*\|\s*-\s*$|\s+-\s+-\s*$|\|\s*0\s*\|\s*0\b)", txt):
            is_explicit_zero = True

        if is_explicit_zero:
            continue

        for seg in segments:
            # Remove note references like "(Nota 4)", "(Anexo I)", or trailing numbers in plain PDF lines
            clean_seg = re.sub(r"\(Notas?[^)]*\)|\(Anexos?[^)]*\)|\(Cuadros?[^)]*\)", "", seg, flags=re.IGNORECASE)
            clean_seg = re.sub(r"[\s$]+(?:\(\s*)?-?\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?(?:\s*\))?.*$", "", clean_seg)
            clean_seg = re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", clean_seg.strip())
            clean_seg = clean_seg.strip(" \t-|:$0123456789.,;")
            if "(" in clean_seg and not clean_seg.endswith(")"):
                clean_seg = clean_seg + ")"

            # Skip narrative sentences or headers
            if len(clean_seg) < 4 or len(clean_seg) > 78 or len(clean_seg.split()) > 10:
                continue

            for cat in RUBRO_POLICY_CATALOG:
                if cat["id"] in seen_ids:
                    continue
                for pat in cat["patterns"]:
                    if re.search(pat, clean_seg, re.IGNORECASE):
                        exact_title = clean_seg[0].upper() + clean_seg[1:]
                        is_leg = any(kw in exact_title.lower() for kw in LEGACY_RUBRO_KEYWORDS)
                        found_rubros.append({
                            "id": cat["id"],
                            "order": cat["order"],
                            "section": cat["section"],
                            "norm_ref": cat["norm_ref"],
                            "exact_name_in_balance": exact_title,
                            "is_legacy_name": is_leg,
                            "default_name": cat["default_name"],
                            "policy_ep": cat["policy_ep"],
                            "policy_re": cat["policy_re"]
                        })
                        seen_ids.add(cat["id"])
                        break

    # Fallback: if the uploaded document is a cover page or sparse scan, include standard core rubros
    if len(found_rubros) < 3:
        core_defaults = [
            "caja_bancos",
            "inversiones_financieras",
            "cuentas_por_cobrar",
            "creditos_impositivos",
            "otras_cuentas_por_cobrar",
            "bienes_uso",
            "proveedores_deudas_comerciales",
            "deudas_laborales",
            "deudas_fiscales",
            "otras_deudas_anticipos_fondos",
            "previsiones_pasivo"
        ]
        for cat in RUBRO_POLICY_CATALOG:
            if cat["id"] in seen_ids or cat["id"] not in core_defaults:
                continue
            found_rubros.append({
                "id": cat["id"],
                "order": cat["order"],
                "section": cat["section"],
                "norm_ref": cat["norm_ref"],
                "exact_name_in_balance": cat["default_name"].split(" / ")[0],
                "is_legacy_name": False,
                "default_name": cat["default_name"],
                "policy_ep": cat["policy_ep"],
                "policy_re": cat["policy_re"]
            })
            seen_ids.add(cat["id"])

    # Always sort by Balance Sheet order (Activo Corriente/No Corriente -> Pasivo Corriente/No Corriente)
    found_rubros.sort(key=lambda x: x["order"])
    return found_rubros


def generate_tailored_notes_package(
    extracted_doc: Dict[str, Any],
    validation_result: Dict[str, Any],
    output_pdf_path: str,
    output_docx_path: str
) -> Dict[str, Any]:
    """
    Generates the complete, standalone 'ANEXO TEXTUAL DE NOTAS A LOS ESTADOS CONTABLES (RT 54 - T.O. RT 59)'
    custom-tailored to the exact rubric names present in the uploaded balance, ready to attach
    to the user's existing numeric Note tables (Cuadros de Notas Activo y Pasivo).
    """
    meta = validation_result["entity_metadata"]
    numbers = validation_result["financial_summary"]
    lines = extracted_doc["lines"]
    full_text = extracted_doc["full_text"]

    entity_name = meta.get("denominacion", "ENTIDAD")
    if "No detectada" in entity_name or "DENOMINACIÓN" in entity_name.upper():
        entity_name = "[COMPLETAR DENOMINACIÓN DE LA ENTIDAD]"

    cuit = meta.get("cuit", "XX-XXXXXXXX-X")
    if "No detectada" in cuit or "Plantilla" in cuit:
        cuit = "[COMPLETAR CUIT]"

    closing_date = meta.get("fecha_cierre", "dd/mm/aaaa")
    if "No detectada" in closing_date:
        closing_date = "31/12/2025"

    is_esfl = meta.get("es_sin_fines_de_lucro", False)
    cat_rt54 = meta.get("categoria_rt54", "Entidad Pequeña (EP)")
    if "No especificada" in cat_rt54:
        cat_rt54 = "Entidad Pequeña (EP)"
    is_ep = "Pequeña" in cat_rt54 or "EP" in cat_rt54

    # Detect the exact rubros used in the user's balance
    active_rubros = detect_active_rubros_from_balance(lines, full_text)

    # Format fecha de inicio del ejercicio in words (e.g. "01 de julio de 2025")
    m_ini = re.search(r"iniciado\s+el\s+(\d{1,2}[/-]\d{1,2}[/-]20\d\d)", full_text, re.IGNORECASE)
    if m_ini:
        fecha_inicio_raw = m_ini.group(1).replace("-", "/")
    else:
        c_parts = closing_date.split("/")
        if len(c_parts) == 3:
            cd, cm, cy = int(c_parts[0]), int(c_parts[1]), int(c_parts[2])
            if cm == 12 and cd == 31:
                fecha_inicio_raw = f"01/01/{cy}"
            elif cm == 6 and cd == 30:
                fecha_inicio_raw = f"01/07/{cy - 1}"
            else:
                fecha_inicio_raw = f"01/{((cm % 12) + 1):02d}/{cy - 1 if cm != 12 else cy}"
        else:
            fecha_inicio_raw = "01/07/2025"

    months_sp = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    parts_ini = fecha_inicio_raw.split("/")
    if len(parts_ini) == 3 and 1 <= int(parts_ini[1]) <= 12:
        fecha_inicio_txt = f"{int(parts_ini[0]):02d} de {months_sp[int(parts_ini[1]) - 1]} de {parts_ini[2]}"
    else:
        fecha_inicio_txt = fecha_inicio_raw

    efe_clause = (
        "• El Estado de Flujos de Efectivo ha sido preparado por el método directo en forma abierta (conforme al párrafo 666 de la RT 54 aplicable a entidades sin fines de lucro)."
        if is_esfl else
        "• El Estado de Flujos de Efectivo ha sido presentado por el método directo, mostrando la información ajustada por inflación en forma sintética (párrafo 666 de la RT 54)."
    )

    caja_act = fmt_ars(numbers.get("caja_y_bancos", {}).get("actual") or numbers.get("efe_efectivo_cierre", {}).get("actual"))
    caja_cmp = fmt_ars(numbers.get("caja_y_bancos", {}).get("comparativo") or numbers.get("efe_efectivo_cierre", {}).get("comparativo"))
    if "No detectado" in caja_act:
        caja_concil_str = "concilia directamente con el saldo expuesto en el rubro 'Caja y bancos' del Estado de Situación Patrimonial."
    else:
        caja_concil_str = f"concilia directamente con el rubro 'Caja y bancos' del Estado de Situación Patrimonial (Ejercicio Actual: {caja_act} | Ejercicio Comparativo: {caja_cmp})."

    # Build Section 1: NOTA 1 - BASES DE PREPARACIÓN (Modelo Oficial CPCECABA - Res. P. 460/2024)
    nota_1_items = [
        {
            "sub": "1.1. Bases de preparación de los estados contables",
            "text": (
                f"Los presentes estados contables están expresados en pesos en moneda homogénea de cierre (Nota 1.3) y han sido preparados de conformidad con la "
                f"Resolución Técnica N° 54 - T.O. RT 59 -, (RT 54), emitida por la Federación Argentina de Consejos Profesionales de Ciencias Económicas (FACPCE), "
                f"adoptada por la Resolución P. N° 460/2024 del Consejo Profesional de Ciencias Económicas de la Ciudad Autónoma de Buenos Aires (CPCECABA).\n\n"
                f"Esta resolución fue aplicada por la entidad a partir del {fecha_inicio_txt}, utilizando el enfoque retroactivo simplificado, "
                f"por lo que la entidad ha aplicado los cambios de la norma retroactivamente sólo readecuando las cifras comparativas incluidas en el estado de situación patrimonial y notas y anexos relacionados.\n\n"
                f"Esta resolución fue aplicada por la entidad sin generar efectos en el reconocimiento y medición de los activos, pasivos y patrimonio neto, "
                f"excepto por las adecuaciones en las denominaciones y exposición de ciertos rubros.\n\n"
                f"Adicionalmente, se detallan las principales reclasificaciones de presentación originadas por la aplicación de la RT 54 en el Estado de situación patrimonial:\n"
                f"a) El saldo correspondiente a créditos impositivos que se presentaba en el rubro 'Otros créditos' del Estado de situación patrimonial, se reclasificó al rubro 'Créditos Impositivos'.\n"
                f"b) Las deudas comerciales con proveedores de bienes y servicios se presentan bajo la denominación 'Proveedores de bienes y servicios' (o 'Deudas con proveedores de bienes y servicios').\n"
                f"c) Las obligaciones con el personal y de la seguridad social se agrupan en el rubro 'Deudas laborales y previsionales'.\n"
                f"d) Las colocaciones a corto plazo se exponen bajo la denominación 'Inversiones financieras'."
            )
        },
        {
            "sub": "1.2. Clasificación de la entidad (RT 54 párr. 6 a 14)",
            "text": f"De acuerdo con lo establecido por la referida norma y los parámetros de ingresos vigentes, la Entidad reviste el carácter de {cat_rt54}."
        },
        {
            "sub": "1.3. Unidad de medida (Expresión en Moneda Homogénea - RT 54 párr. 176 a 200)",
            "text": (
                f"Los presentes estados contables han sido preparados en moneda homogénea a la fecha de cierre del ejercicio ({closing_date}), reconociendo en forma integral los efectos de la inflación de conformidad con lo establecido en los párrafos 176 a 200 de la RT 54.\n"
                "Con fines comparativos, se incluyen cifras patrimoniales, de resultados (o recursos y gastos), de evolución del patrimonio neto y de flujos de efectivo correspondientes al ejercicio anterior, íntegramente reexpresadas en moneda de poder adquisitivo de cierre del presente ejercicio utilizando la serie de índices de precios publicada por la FACPCE.\n"
                "Los resultados derivados del mantenimiento de activos y pasivos expuestos a los cambios en el poder adquisitivo de la moneda a lo largo del ejercicio se exponen unificados en la línea «Resultados financieros y por tenencia (incluyendo el RECPAM)».\n"
                "Para la preparación de los presentes estados contables, el órgano de administración ha utilizado las siguientes simplificaciones dispuestas por la RT 54:\n"
                "• No ha presentado la nota requerida por el párrafo 659 de la RT 54 ni la del inciso f) del párrafo 599.\n"
                f"{efe_clause}\n"
                "• Optó por presentar en una sola línea los resultados financieros y por tenencia, incluyendo al resultado por exposición al cambio en el poder adquisitivo de la moneda (RECPAM).\n"
                "• De conformidad con el principio de síntesis y flexibilidad (RT 54 párr. 55 y 56 inc. a), se han suprimido los rubros, filas y anexos que no presentan saldo en el ejercicio actual ni en el comparativo.\n"
                "Los presentes estados contables deben ser leídos e interpretados considerando las limitaciones que la utilización de las dispensas antes mencionadas podría provocar sobre la información en ellos contenida."
            )
        },
        {
            "sub": "1.4. Uso de estimaciones contables",
            "text": (
                "La preparación de los estados contables requiere que el órgano de administración realice estimaciones y evaluaciones que afectan el monto de los activos y pasivos registrados y las contingencias reveladas a la fecha de cierre, así como los ingresos y egresos del ejercicio. Los resultados reales futuros pueden diferir de las estimaciones realizadas a la fecha de preparación de los presentes estados contables."
            )
        },
        {
            "sub": "1.5. Información comparativa",
            "text": (
                "Las cifras correspondientes al ejercicio económico anterior se presentan únicamente con fines comparativos y han sido reexpresadas en moneda homogénea de cierre del presente ejercicio conforme a lo indicado en la Nota 1.3, sin que dicho ajuste modifique las decisiones tomadas sobre la base de los estados contables del ejercicio anterior."
            )
        }
    ]

    # Build Section 2: NOTA 2 - POLÍTICAS CONTABLES SOBRE CRITERIOS DE MEDICIÓN (Respetando los nombres exactos de rubros del balance subido)
    nota_2_items = [
        {
            "sub": "2.1. Activos y pasivos en moneda extranjera (de corresponder)",
            "norm_ref": "RT 54 párr. 201-206",
            "text": "Los activos y pasivos en moneda extranjera se convierten al tipo de cambio comprador o vendedor, respectivamente, vigente a la fecha de cierre de los estados contables."
        }
    ]

    for idx, r in enumerate(active_rubros, start=2):
        policy_txt = r["policy_ep"] if is_ep else r["policy_re"]
        nota_2_items.append({
            "sub": f"2.{idx}. {r['exact_name_in_balance']}",
            "norm_ref": r.get("norm_ref", "RT 54"),
            "text": policy_txt
        })

    next_idx = len(nota_2_items) + 1
    nota_2_items.append({
        "sub": f"2.{next_idx}. Valor recuperable de Bienes de Uso, Propiedades de Inversión y Activos Intangibles (RT 54 párr. 150-165)",
        "norm_ref": "RT 54 párr. 150-165",
        "text": (
            "La Entidad evalúa en cada cierre la existencia de indicios de deterioro que puedan determinar la necesidad de comparar la medición contable de los activos no corrientes con su valor recuperable (considerando el requisito de tres ejercicios consecutivos con resultado positivo para la dispensa en Entidades Pequeñas). En el presente ejercicio no se han identificado indicios de desvalorización y los valores registrados no superan su valor recuperable."
        )
    })

    if is_esfl:
        nota_2_items.append({
            "sub": f"2.{next_idx + 1}. Situación frente al Impuesto a las Ganancias y Recursos / Gastos Ordinarios",
            "norm_ref": "RT 54 párr. 719-741",
            "text": (
                "La Entidad reviste el carácter de sujeto exento en el Impuesto a las Ganancias en virtud del artículo 26 de la Ley de Impuesto a las Ganancias, por lo que no corresponde contabilizar cargo ni pasivo por dicho tributo. Los recursos y gastos para fines generales y específicos se reconocen en función de su devengamiento y se reexpresan en moneda de cierre conforme a la Nota 1.3."
            )
        })
        nota_2_items.append({
            "sub": f"2.{next_idx + 2}. Cuentas del Patrimonio Neto (Capital Social / Fondos Sociales)",
            "norm_ref": "RT 54 párr. 508-524",
            "text": (
                "Se encuentran reexpresadas de acuerdo con lo mencionado en la Nota 1.3. El Capital Social / Aporte de los Asociados se mantiene a su valor nominal y el ajuste derivado de su reexpresión en moneda homogénea se expone en la cuenta «Ajuste del capital»."
            )
        })
    else:
        nota_2_items.append({
            "sub": f"2.{next_idx + 1}. Costo de los bienes vendidos / servicios prestados",
            "norm_ref": "RT 54 párr. 525-543",
            "text": "Se calcula aplicando la técnica de diferencia de inventario (o acumulación de costos de servicios prestados) sobre valores reexpresados en moneda homogénea de cierre de acuerdo con la Nota 1.3."
        })
        nota_2_items.append({
            "sub": f"2.{next_idx + 2}. Impuesto a las Ganancias",
            "norm_ref": "RT 54 párr. 571-600",
            "text": (
                "El cargo contable por Impuesto a las Ganancias se reconoce en función del impuesto determinado del ejercicio según la legislación fiscal vigente (haciendo uso de la opción para Entidades Pequeñas de no presentar la nota del inciso f del párrafo 599 de la RT 54)."
                if is_ep else
                "La Entidad reconoce el cargo por Impuesto a las Ganancias aplicando el método del impuesto diferido, considerando las diferencias temporarias entre las mediciones contables e impositivas de activos y pasivos."
            )
        })
        nota_2_items.append({
            "sub": f"2.{next_idx + 3}. Cuentas de Capital Social y Patrimonio Neto (RG IGJ 15/2024 Art. 240)",
            "norm_ref": "RT 54 párr. 508-524",
            "text": (
                "Se encuentran reexpresadas de acuerdo con lo mencionado en la Nota 1.3. Las cuentas «Capital suscripto» y, de existir, «Aportes irrevocables» se mantienen a sus valores nominales y el ajuste derivado de su reexpresión monetaria se expone en «Ajuste del capital» y «Ajuste de Aportes Irrevocables» (RG IGJ N° 15/2024 Art. 240)."
            )
        })

    # Build Section 3: NOTA 3 - VINCULACIÓN CON LOS CUADROS DE DETALLE DE NOTAS DE ACTIVO Y PASIVO Y CONCILIACIÓN DEL EFECTIVO
    rubros_list_str = ", ".join(f"«{r['exact_name_in_balance']}»" for r in active_rubros) if active_rubros else "los rubros con saldo del Activo y del Pasivo"
    nota_3_items = [
        {
            "sub": "3.1. Remisión a los Cuadros de Composición de Rubros del Activo y del Pasivo",
            "text": (
                f"El detalle cuantitativo y la apertura en Corriente y No Corriente de los rubros que integran el Estado de Situación Patrimonial al {closing_date} y al cierre comparativo "
                f"({rubros_list_str}) se exponen en los «Cuadros de Detalle de Notas de Activo y de Pasivo» y Anexos que acompañan y forman parte integrante de la presente información complementaria."
            )
        },
        {
            "sub": "3.2. Conciliación del Efectivo y sus Equivalentes (Estado de Flujos de Efectivo vs. Estado de Situación Patrimonial - RT 54 párr. 672 inc. a)",
            "text": (
                f"En cumplimiento de lo dispuesto por el párrafo 672 inciso a) de la RT 54 y los modelos del CPCECABA, se informa que el concepto «Efectivo y sus equivalentes al cierre del ejercicio» considerado en el Estado de Flujos de Efectivo {caja_concil_str}"
            )
        },
        {
            "sub": "3.3. Bienes de disponibilidad restringida, gravámenes y contingencias posibles (RT 54 párr. 499)",
            "text": (
                f"Al {closing_date} no existen bienes del activo cuya disponibilidad se encuentre restringida, ni activos gravados con hipoteca, prenda u otros derechos reales, ni existen situaciones contingentes posibles (no probables) pendientes de revelación en los presentes estados contables."
            )
        },
        {
            "sub": "3.4. Hechos posteriores al cierre del ejercicio",
            "text": (
                f"Entre la fecha de cierre del ejercicio ({closing_date}) y la fecha de emisión de los presentes estados contables no se han producido hechos ni circunstancias significativas de carácter financiero o económico que afecten sustancialmente la situación patrimonial, los resultados o la evolución del efectivo de la Entidad y que requieran ajuste o revelación adicional."
            )
        }
    ]

    # Render standalone PDF and Editable Word (.docx)
    _render_notas_pdf(
        output_pdf_path,
        entity_name=entity_name,
        cuit=cuit,
        closing_date=closing_date,
        nota_1_items=nota_1_items,
        nota_2_items=nota_2_items,
        nota_3_items=nota_3_items,
        active_rubros=active_rubros
    )

    _render_notas_docx(
        output_docx_path,
        entity_name=entity_name,
        cuit=cuit,
        closing_date=closing_date,
        nota_1_items=nota_1_items,
        nota_2_items=nota_2_items,
        nota_3_items=nota_3_items,
        active_rubros=active_rubros
    )

    uses_legacy_names = any(r.get("is_legacy_name", False) for r in active_rubros)
    return {
        "entity_name": entity_name,
        "cuit": cuit,
        "closing_date": closing_date,
        "categoria_rt54": cat_rt54,
        "uses_legacy_rubro_names": uses_legacy_names,
        "detected_rubros": [
            {
                "id": r["id"],
                "exact_name_in_balance": r["exact_name_in_balance"],
                "section": r.get("section", "Activo/Pasivo"),
                "norm_ref": r.get("norm_ref", "RT 54"),
                "is_legacy_name": r.get("is_legacy_name", False)
            }
            for r in active_rubros
        ],
        "nota_1_bases": [{"subtitulo": x["sub"], "texto": x["text"]} for x in nota_1_items],
        "nota_2_politicas": [{"subtitulo": x["sub"], "norm_ref": x.get("norm_ref", "RT 54"), "texto": x["text"]} for x in nota_2_items],
        "nota_3_remision_cuadros": [{"subtitulo": x["sub"], "texto": x["text"]} for x in nota_3_items]
    }


def _render_notas_pdf(
    output_pdf_path: str,
    entity_name: str,
    cuit: str,
    closing_date: str,
    nota_1_items: List[Dict[str, str]],
    nota_2_items: List[Dict[str, str]],
    nota_3_items: List[Dict[str, str]],
    active_rubros: List[Dict[str, Any]]
):
    os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)
    doc = SimpleDocTemplate(
        output_pdf_path,
        pagesize=A4,
        leftMargin=18*mm,
        rightMargin=18*mm,
        topMargin=18*mm,
        bottomMargin=18*mm
    )
    w_content = A4[0] - 36*mm
    story = []

    h1_style = ParagraphStyle('H1N', fontName=FONT_BOLD, fontSize=11.5, leading=14.5, textColor=NAVY_PRIMARY, alignment=TA_CENTER)
    sub_center = ParagraphStyle('SubCN', fontName=FONT_REG, fontSize=8.5, leading=11, textColor=TEXT_DARK, alignment=TA_CENTER)
    sec_bar_style = ParagraphStyle('SBN', fontName=FONT_BOLD, fontSize=8.5, leading=11, textColor=WHITE)
    sub_head = ParagraphStyle('SHN', fontName=FONT_BOLD, fontSize=8.4, leading=11, textColor=NAVY_PRIMARY)
    body = ParagraphStyle('BN', fontName=FONT_REG, fontSize=8.0, leading=11.2, textColor=TEXT_DARK, alignment=TA_JUSTIFY)

    def make_bar(title: str):
        t = Table([[Paragraph(title.upper(), sec_bar_style)]], colWidths=[w_content])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), NAVY_PRIMARY),
            ('TOPPADDING', (0, 0), (-1, -1), 1.8*mm),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1.8*mm),
            ('LEFTPADDING', (0, 0), (-1, -1), 3*mm),
        ]))
        return t

    story.append(Paragraph(entity_name.upper(), h1_style))
    story.append(Spacer(1, 1*mm))
    story.append(Paragraph("NOTAS A LOS ESTADOS CONTABLES (ANEXO TEXTUAL DE POLÍTICAS Y BASES RT 54 - T.O. RT 59)", sub_center))
    story.append(Paragraph(f"Por el ejercicio finalizado el {closing_date}, comparativo con el ejercicio anterior — En moneda homogénea", ParagraphStyle('ItC', fontName=FONT_ITA, fontSize=7.8, leading=10, textColor=TEXT_MUTED, alignment=TA_CENTER)))
    story.append(Spacer(1, 4*mm))

    # Sección 1
    story.append(make_bar("1. Notas Generales y Bases de Preparación de los Estados Contables (RT 54)"))
    story.append(Spacer(1, 2.5*mm))
    for item in nota_1_items:
        story.append(Paragraph(item["sub"], sub_head))
        story.append(Spacer(1, 0.8*mm))
        for p in item["text"].split("\n"):
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 1.2*mm))
        story.append(Spacer(1, 1.2*mm))

    # Sección 2
    story.append(make_bar("2. Políticas Contables sobre Criterios de Medición (Rubros con Saldo en el Balance)"))
    story.append(Spacer(1, 2.5*mm))
    for item in nota_2_items:
        story.append(Paragraph(item["sub"], sub_head))
        story.append(Spacer(1, 0.8*mm))
        for p in item["text"].split("\n"):
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 1.2*mm))
        story.append(Spacer(1, 1.2*mm))

    # Sección 3
    story.append(make_bar("3. Remisión a Cuadros de Activo/Pasivo, Conciliación del Efectivo y Otras Revelaciones"))
    story.append(Spacer(1, 2.5*mm))
    for item in nota_3_items:
        story.append(Paragraph(item["sub"], sub_head))
        story.append(Spacer(1, 0.8*mm))
        for p in item["text"].split("\n"):
            story.append(Paragraph(p, body))
            story.append(Spacer(1, 1.2*mm))
        story.append(Spacer(1, 1.2*mm))

    doc.build(
        story,
        canvasmaker=lambda *args, **kwargs: ConsejoSequentialCanvas(
            *args, entity_name=entity_name, cuit=cuit, closing_date=closing_date, **kwargs
        )
    )


def _render_notas_docx(
    output_docx_path: str,
    entity_name: str,
    cuit: str,
    closing_date: str,
    nota_1_items: List[Dict[str, str]],
    nota_2_items: List[Dict[str, str]],
    nota_3_items: List[Dict[str, str]],
    active_rubros: List[Dict[str, Any]]
):
    os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
    d = docx.Document()

    p_top = d.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_top = p_top.add_run(f"{entity_name.upper()}\nNOTAS A LOS ESTADOS CONTABLES\nPor el ejercicio finalizado el {closing_date} comparativo en pesos con el ejercicio anterior. En moneda homogénea.")
    r_top.bold = True
    r_top.font.size = Pt(11.5)
    r_top.font.color.rgb = RGBColor(13, 31, 60)

    d.add_heading("1. NOTAS GENERALES Y BASES DE PREPARACIÓN (RT 54 - T.O. RT 59)", level=2)
    for item in nota_1_items:
        p_h = d.add_paragraph()
        r_h = p_h.add_run(item["sub"])
        r_h.bold = True
        for line in item["text"].split("\n"):
            p_b = d.add_paragraph(line)
            p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    d.add_heading("2. POLÍTICAS CONTABLES SOBRE CRITERIOS DE MEDICIÓN (SEGÚN RUBROS DEL BALANCE)", level=2)
    for item in nota_2_items:
        p_h = d.add_paragraph()
        r_h = p_h.add_run(item["sub"])
        r_h.bold = True
        for line in item["text"].split("\n"):
            p_b = d.add_paragraph(line)
            p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    d.add_heading("3. COMPOSICIÓN DE RUBROS, CONCILIACIÓN DEL EFECTIVO Y OTRAS REVELACIONES", level=2)
    for item in nota_3_items:
        p_h = d.add_paragraph()
        r_h = p_h.add_run(item["sub"])
        r_h.bold = True
        for line in item["text"].split("\n"):
            p_b = d.add_paragraph(line)
            p_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    d.save(output_docx_path)
